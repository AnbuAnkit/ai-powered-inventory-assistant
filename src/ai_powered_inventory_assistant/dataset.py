"""SQLite storage and query helpers for the sales workbook."""

from __future__ import annotations

import sqlite3
from contextlib import closing
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

EXPECTED_HEADERS = (
    "Order_ID",
    "Date",
    "Region",
    "Salesperson",
    "Product",
    "Category",
    "Quantity",
    "Sales",
    "Profit",
)


def initialize_database(database_path: Path, workbook_path: Path) -> int:
    """Import the workbook's Sales_Data sheet into SQLite and return row count."""
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Sales workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Sales_Data" not in workbook.sheetnames:
            raise ValueError("Sales workbook must contain a 'Sales_Data' sheet.")

        worksheet = workbook["Sales_Data"]
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if (
            headers is None
            or tuple(headers[: len(EXPECTED_HEADERS)]) != EXPECTED_HEADERS
        ):
            raise ValueError("Sales_Data has missing or unexpected column headers.")

        records: list[tuple[Any, ...]] = []
        for row_number, values in enumerate(rows, start=2):
            if not values or all(value is None for value in values):
                continue
            if len(values) < len(EXPECTED_HEADERS):
                raise ValueError(f"Sales_Data row {row_number} has too few columns.")
            (
                order_id,
                sale_date,
                region,
                salesperson,
                product,
                category,
                quantity,
                sales,
                profit,
            ) = values[: len(EXPECTED_HEADERS)]
            required_values = (
                order_id,
                sale_date,
                region,
                salesperson,
                product,
                category,
                quantity,
                sales,
                profit,
            )
            if any(value is None for value in required_values):
                raise ValueError(f"Sales_Data row {row_number} contains an empty required value.")
            if isinstance(sale_date, datetime):
                sale_date = sale_date.date()
            if isinstance(sale_date, date):
                sale_date = sale_date.isoformat()
            else:
                sale_date = str(sale_date)
            records.append(
                (
                    int(order_id),
                    sale_date,
                    str(region),
                    str(salesperson),
                    str(product),
                    str(category),
                    int(quantity),
                    float(sales),
                    float(profit),
                )
            )
    finally:
        workbook.close()

    database_path.parent.mkdir(parents=True, exist_ok=True)
    with closing(sqlite3.connect(database_path)) as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS sales (
                order_id INTEGER PRIMARY KEY,
                sale_date TEXT NOT NULL,
                region TEXT NOT NULL,
                salesperson TEXT NOT NULL,
                product TEXT NOT NULL,
                category TEXT NOT NULL,
                quantity INTEGER NOT NULL,
                sales REAL NOT NULL,
                profit REAL NOT NULL
            )
            """
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS idx_sales_region ON sales(region)"
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS idx_sales_category ON sales(category)"
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS inventory_items (
                item_id INTEGER PRIMARY KEY,
                product TEXT NOT NULL COLLATE NOCASE UNIQUE,
                stock_quantity INTEGER NOT NULL CHECK (stock_quantity >= 0),
                supplier TEXT NOT NULL DEFAULT '',
                reorder_point INTEGER NOT NULL DEFAULT 0 CHECK (reorder_point >= 0),
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        with connection:
            connection.execute("DELETE FROM sales")
            connection.executemany(
                """
                INSERT INTO sales (
                    order_id, sale_date, region, salesperson, product, category,
                    quantity, sales, profit
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                records,
            )
    return len(records)


def list_inventory_items(
    database_path: Path, *, product: str | None = None, limit: int = 100
) -> list[dict[str, Any]]:
    """Return inventory items, optionally matching a product name."""
    query = """
        SELECT product, stock_quantity, supplier, reorder_point, updated_at
        FROM inventory_items
    """
    parameters: list[Any] = []
    if product is not None:
        query += " WHERE product LIKE ? COLLATE NOCASE"
        parameters.append(f"%{product}%")
    query += " ORDER BY product LIMIT ?"
    parameters.append(limit)
    with closing(sqlite3.connect(database_path)) as connection:
        connection.row_factory = sqlite3.Row
        rows = connection.execute(query, parameters).fetchall()
    return [dict(row) for row in rows]


def add_inventory_item(
    database_path: Path,
    *,
    product: str,
    stock_quantity: int,
    supplier: str = "",
    reorder_point: int = 0,
) -> dict[str, Any]:
    """Create an inventory item and return the stored row."""
    product = product.strip()
    if not product:
        raise ValueError("product must not be empty")
    if stock_quantity < 0 or reorder_point < 0:
        raise ValueError("stock_quantity and reorder_point must not be negative")
    with closing(sqlite3.connect(database_path)) as connection:
        with connection:
            connection.execute(
                """
                INSERT INTO inventory_items (
                    product, stock_quantity, supplier, reorder_point
                ) VALUES (?, ?, ?, ?)
                """,
                (product, stock_quantity, supplier.strip(), reorder_point),
            )
        connection.row_factory = sqlite3.Row
        row = connection.execute(
            """
            SELECT product, stock_quantity, supplier, reorder_point, updated_at
            FROM inventory_items WHERE product = ? COLLATE NOCASE
            """,
            (product,),
        ).fetchone()
    if row is None:
        raise RuntimeError("Inserted inventory item could not be read back")
    return dict(row)


def update_inventory_item(
    database_path: Path,
    *,
    product: str,
    stock_quantity: int | None = None,
    supplier: str | None = None,
    reorder_point: int | None = None,
) -> dict[str, Any]:
    """Update provided fields on an existing inventory item."""
    updates: list[str] = []
    values: list[Any] = []
    if stock_quantity is not None:
        if stock_quantity < 0:
            raise ValueError("stock_quantity must not be negative")
        updates.append("stock_quantity = ?")
        values.append(stock_quantity)
    if supplier is not None:
        updates.append("supplier = ?")
        values.append(supplier.strip())
    if reorder_point is not None:
        if reorder_point < 0:
            raise ValueError("reorder_point must not be negative")
        updates.append("reorder_point = ?")
        values.append(reorder_point)
    if not updates:
        raise ValueError("At least one inventory field must be provided")

    product = product.strip()
    if not product:
        raise ValueError("product must not be empty")
    values.append(product)
    with closing(sqlite3.connect(database_path)) as connection:
        with connection:
            cursor = connection.execute(
                f"""
                UPDATE inventory_items
                SET {", ".join(updates)}, updated_at = CURRENT_TIMESTAMP
                WHERE product = ? COLLATE NOCASE
                """,
                values,
            )
            if cursor.rowcount == 0:
                raise ValueError(f"No inventory item found for product '{product}'")
        connection.row_factory = sqlite3.Row
        row = connection.execute(
            """
            SELECT product, stock_quantity, supplier, reorder_point, updated_at
            FROM inventory_items WHERE product = ? COLLATE NOCASE
            """,
            (product,),
        ).fetchone()
    if row is None:
        raise RuntimeError("Updated inventory item could not be read back")
    return dict(row)


def delete_inventory_item(database_path: Path, *, product: str) -> bool:
    """Delete an inventory item by product name."""
    product = product.strip()
    if not product:
        raise ValueError("product must not be empty")
    with closing(sqlite3.connect(database_path)) as connection:
        with connection:
            cursor = connection.execute(
                "DELETE FROM inventory_items WHERE product = ? COLLATE NOCASE",
                (product,),
            )
    if cursor.rowcount == 0:
        raise ValueError(f"No inventory item found for product '{product}'")
    return True


def _filters(
    region: str | None,
    category: str | None,
    product: str | None,
    salesperson: str | None,
    start_date: date | None,
    end_date: date | None,
) -> tuple[str, list[str]]:
    conditions: list[str] = []
    parameters: list[str] = []
    for column, value in (
        ("region", region),
        ("category", category),
        ("product", product),
        ("salesperson", salesperson),
    ):
        if value is not None:
            conditions.append(f"{column} = ? COLLATE NOCASE")
            parameters.append(value)
    if start_date is not None:
        conditions.append("sale_date >= ?")
        parameters.append(start_date.isoformat())
    if end_date is not None:
        conditions.append("sale_date <= ?")
        parameters.append(end_date.isoformat())
    return (" WHERE " + " AND ".join(conditions) if conditions else ""), parameters


def query_sales(
    database_path: Path,
    *,
    region: str | None = None,
    category: str | None = None,
    product: str | None = None,
    salesperson: str | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
    limit: int = 50,
    offset: int = 0,
) -> dict[str, Any]:
    """Return filtered, paginated sales records and the matching total."""
    where, parameters = _filters(
        region, category, product, salesperson, start_date, end_date
    )
    with closing(sqlite3.connect(database_path)) as connection:
        connection.row_factory = sqlite3.Row
        total = connection.execute(
            "SELECT COUNT(*) FROM sales" + where, parameters
        ).fetchone()[0]
        rows = connection.execute(
            """
            SELECT order_id, sale_date AS date, region, salesperson, product,
                   category, quantity, sales, profit
            FROM sales
            """
            + where
            + " ORDER BY order_id LIMIT ? OFFSET ?",
            [*parameters, limit, offset],
        ).fetchall()
    return {
        "total": total,
        "limit": limit,
        "offset": offset,
        "rows": [dict(row) for row in rows],
    }


def summarize_sales(
    database_path: Path,
    *,
    region: str | None = None,
    category: str | None = None,
    product: str | None = None,
    salesperson: str | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
) -> dict[str, Any]:
    """Return totals and regional/category breakdowns for matching sales."""
    where, parameters = _filters(
        region, category, product, salesperson, start_date, end_date
    )
    with closing(sqlite3.connect(database_path)) as connection:
        totals = connection.execute(
            f"""
            SELECT COUNT(*) AS rows, COALESCE(SUM(quantity), 0) AS quantity,
                   COALESCE(SUM(sales), 0) AS sales,
                   COALESCE(SUM(profit), 0) AS profit
            FROM sales{where}
            """,
            parameters,
        ).fetchone()
        breakdowns = {}
        for field in ("region", "category"):
            breakdowns[field] = [
                {
                    field: row[0],
                    "rows": row[1],
                    "sales": row[2],
                    "profit": row[3],
                    "quantity": row[4],
                }
                for row in connection.execute(
                    f"""
                    SELECT {field}, COUNT(*), SUM(sales), SUM(profit), SUM(quantity)
                    FROM sales{where}
                    GROUP BY {field} ORDER BY {field}
                    """,
                    parameters,
                )
            ]
    return {
        "rows": totals[0],
        "quantity": totals[1],
        "sales": totals[2],
        "profit": totals[3],
        **breakdowns,
    }


def rank_products_by_profit(
    database_path: Path,
    *,
    region: str | None = None,
    category: str | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
    limit: int = 10,
) -> list[dict[str, Any]]:
    """Rank products by summed sales profit, highest first."""
    if not 1 <= limit <= 100:
        raise ValueError("limit must be between 1 and 100")
    where, parameters = _filters(
        region, category, None, None, start_date, end_date
    )
    with closing(sqlite3.connect(database_path)) as connection:
        connection.row_factory = sqlite3.Row
        rows = connection.execute(
            """
            SELECT product, SUM(profit) AS total_profit,
                   SUM(sales) AS total_sales, SUM(quantity) AS total_quantity,
                   COUNT(*) AS transaction_count
            FROM sales
            """
            + where
            + """
            GROUP BY product
            ORDER BY total_profit DESC, product
            LIMIT ?
            """,
            [*parameters, limit],
        ).fetchall()
    return [dict(row) for row in rows]
